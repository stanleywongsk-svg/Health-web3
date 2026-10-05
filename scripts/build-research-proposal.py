#!/usr/bin/env python3
"""Complete the supplied COMM7115 PowerPoint template with HealthLoop evidence.

Usage: python scripts/build-research-proposal.py [--template /path/to/template.pptx]
Uses python-pptx 1.0.2. No network, application, or database operations.
"""

from argparse import ArgumentParser
from copy import deepcopy
from datetime import datetime
from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs/proposals/HealthLoop_Business_Proposal_Research_Report.pptx"
# macOS exposes the template's Charter family under this supported family name.
FONT = "Charter"
INK = "163449"
TEAL = "007F78"
MUTED = "5F7D82"

SOURCES = {
    "S1": ("Hong Kong Department of Health, Health Behaviour Survey 2023",
           "https://www.chp.gov.hk/files/pdf/dh_hbs_2023_report_eng.pdf",
           "2023 fieldwork; report published 2025. Executive summary pp. v–vi, printed pp. 22–29. Adult figures are for ages 18+. Population context, not a HealthLoop demand or effect estimate."),
    "S2": ("Apple Support, Use the Health app on your iPhone or iPad",
           "https://support.apple.com/en-hk/104997",
           "Official product description checked 5 October 2026. Does not demonstrate comparative retention."),
    "S3": ("Sweatcoin Terms of Use, sections 7–8",
           "https://sweatco.in/tnc",
           "Official description of movement-based units and a third-party benefits marketplace, checked 5 October 2026. No claim about Hong Kong offer availability or HealthLoop partnerships."),
    "S4": ("Apple App Review Guidelines, 5.1.2(vi) and 5.1.3(i)",
           "https://developer.apple.com/app-store/review/guidelines/#health-and-health-research",
           "Checked 5 October 2026. Platform policy informs product design; it is not legal advice or evidence of approval."),
    "E1": ("HealthLoop implementation brief and requirement register",
           "MEGA_PROMPT.md; docs/requirements.md",
           "Local repository inspected 5 October 2026. Preserves 58 IDs and B/C ownership. Product rules are not medical recommendations."),
    "E2": ("HealthLoop implementation status and executed test evidence",
           "docs/IMPLEMENTATION_STATUS.md; docs/TEST_EVIDENCE.md",
           "Includes dated 24–28 September and 5 October 2026 engineering records. The October tests were run in a separate application session, not this proposal session. Current UI interaction and physical-device acceptance remain unverified."),
    "E3": ("HealthLoop simulator walkthrough and redacted recording",
           "docs/SIMULATOR_WALKTHROUGH_20260928.md; docs/evidence/recordings/HealthLoop-simulator-walkthrough.mp4",
           "28 September 2026. Local test account and real local Auth/Edge/database. No real health readings, awards, merchant redemption or physical-device acceptance."),
    "E4": ("HealthLoop application and business proposal",
           "docs/proposals/HealthLoop_应用说明与商业提案_简体中文.md",
           "28 September 2026. Pricing, costs, cohort thresholds and funding scenarios are unvalidated planning assumptions. This deck proposes the 12-user, 8-task and 2-paid-pilot research details."),
}


# Every value addresses a text shape in the supplied 17-slide template.
# Original slide order, theme, typography, section hierarchy and timings remain.
CONTENT = {
    1: {
        1: "HEALTHLOOP\nBUSINESS PROPOSAL",
        2: "Research report on an iPhone health habit service",
        5: "HealthLoop team · A / B / C / D · Evidence reviewed 5 October 2026",
    },
    2: {
        6: "Hong Kong adults with an iPhone. Initial focus: desk-based workers willing to try a Simplified Chinese app.",
        7: "PAIN HYPOTHESIS",
        8: "A busy day makes it hard to turn activity records into a manageable, repeatable routine.",
        11: "HealthLoop pairs optional health reads with clear daily missions and traceable, noncash points.",
        12: "Evidence: public research and a tested prototype [S1, E2–E3].\nOpen question: will people keep returning without cash rewards?",
    },
    3: {
        1: "Team responsibilities and field choice",
        6: "ASSIGNED RESPONSIBILITY",
        7: "PROJECT CONTRIBUTION",
        9: "A · Product / business",
        10: "Commercial and scope coordination",
        11: "Buyer research, partner evidence and legal coordination",
        13: "B · Mobile",
        14: "iOS experience and native delivery",
        15: "HealthKit adapter and app flow; C reviews technical work",
        17: "C · Backend",
        18: "Data, accounting and security",
        19: "Auth, points and admin controls; B reviews technical work",
        21: "D · Research / QA",
        22: "Design, research and QA support",
        23: "Recruitment, synthesis, usability and black-box QA support",
        30: "Daily habit support: build on a working prototype and test customer demand.",
    },
    4: {
        6: "Adults fitting activity around desk work and commuting. iPhone access and language fit are recruitment criteria, not proven demand.",
        8: "Build a manageable routine and understand progress without adding much manual tracking. This job is a hypothesis to test.",
        11: "Apple Health already records steps [S2]. Interviews will test use of existing trackers, reminders, walking groups or no tool.",
        13: "Public data supports investigating daily activity [S1]. The time, effort and frustration costs for our segment remain unmeasured.",
    },
    5: {
        8: "What value could a routine add beyond existing activity trackers?",
        9: "Done: official-source review [S1–S4]. Next: 12 contextual interviews to understand actual workarounds.",
        11: "Can users explain the mission, points and privacy choices unaided?",
        12: "Next: 8 moderated task sessions, then a cohort with 100 matured activations to test repeat use.",
        14: "Will organisations fund delivery without buying health profiles?",
        15: "Next: 5 budget-holder interviews to test a defined service, price, procurement and delivery capacity.",
        17: "Desk review: 5 October 2026 · Fieldwork proposed, not conducted.\nDecision: refine the segment and approve a bounded pilot only when readiness gates pass.",
    },
    6: {
        5: "SAMPLE STATUS",
        6: "External research participants recorded: 0.\nProposed: 12 adults and 5 budget holders; 8 adults also complete usability tasks.",
        7: "PROPOSED RECRUITMENT AND SESSIONS",
        8: "Two community or workplace channels; 30–45 minute interviews and 20 minute task sessions. Recruitment has not begun.",
        11: "A small volunteer sample cannot represent Hong Kong. Include sceptics and consent refusers; test language fit and report exclusions.",
        13: "Plain-language consent, optional recording and participant codes. No raw health uploads. Set withdrawal and deletion procedures before recruitment.",
    },
    7: {
        9: "14.8% insufficient activity;\n33.9% sit or recline ≥8 hours on a typical day.",
        10: "S1 · Hong Kong DH\n2023 survey; 2025 report\nAdults aged 18+",
        11: "A relevant activity context. Self-report data; neither a market-size estimate nor evidence of app demand.",
        13: "Apple Health already counts steps and brings health records together.",
        14: "S2 · Apple Support\nOfficial product guide\nChecked 5 October 2026",
        15: "A strong existing alternative. HealthLoop must add value beyond a second step-count screen.",
        17: "Sweatcoin links verified movement to in-app units and a benefits marketplace.",
        18: "S3 · Sweatcoin\nTerms, sections 7–8\nChecked 5 October 2026",
        19: "Activity incentives already compete for attention. HealthLoop has no signed merchant network to claim.",
    },
    8: {
        1: "Findings from current evidence",
        6: "A local problem context exists; demand is unproven.",
        8: "The public survey supports investigating activity habits. It does not explain our segment’s unmet needs. [S1]",
        11: "Step counting and activity incentives are already served.",
        13: "Apple Health and Sweatcoin are existing alternatives; mission clarity and trust still need testing. [S2–S3]",
        16: "Technical feasibility has progressed; adoption is unmeasured.",
        18: "355 tests logged on 5 October; September’s simulator flow was recorded. Current UI and two-iPhone checks remain open. [E2–E3]",
    },
    9: {
        1: "Provisional persona to test",
        5: "Hong Kong desk-based iPhone user",
        6: "Recruitment hypothesis from the product brief; no interview-derived persona yet. [E1, E4]",
        9: "Keep a manageable activity routine and see what happened today without treating rest as failure.",
        10: "ASSUMED CONTEXT AND BEHAVIOUR",
        11: "Fits activity around work and commuting; already has an iPhone. App language and routine fit require validation.",
        12: "ASSUMED PAIN AND WORKAROUND",
        13: "May check a tracker occasionally yet struggle to return consistently. Interviews must confirm or reject this pattern.",
        15: "Test one clear next action, an honest data state and optional consent before adding more incentives.",
    },
    10: {
        1: "Proposed scenario and customer journey",
        5: "HYPOTHESIS · A desk-based adult wants to restart a walking routine during a workday break.",
        8: "Action\nNotices a break after a long desk session.",
        9: "Test assumption\nDoes time pressure disrupt the routine?",
        12: "Action\nChecks an existing activity record and today’s goal.",
        13: "Test assumption\nIs another app worth opening?",
        16: "Action\nReads a mission; chooses whether to read and sync data.",
        17: "Test assumption\nAre data choices and point rules clear?",
        20: "Action\nSees posted or pending points; decides whether to return.",
        21: "Test assumption\nDoes feedback matter without cash value?",
        23: "PRIORITY TOUCHPOINT",
        24: "Make the first mission and consent consequences understandable; test without coaching.",
    },
    11: {
        1: "Pain hypotheses and opportunity choice",
        5: "PAIN HYPOTHESIS",
        9: "Activity records may not create a repeatable routine.",
        10: "S1–S2 provide context; no interview confirmation yet.",
        11: "Small daily choices and a weekly progress view.",
        13: "Users may misunderstand points or data permissions.",
        14: "E1 rules and E3 walkthrough; comprehension untested.",
        15: "Explain sources, pending states and separate consent.",
        17: "Organisers may lack capacity to run habit programmes.",
        18: "E4 business hypothesis; no buyer interviews recorded.",
        19: "Test a fixed content, onboarding and support package.",
        22: "Prioritise comprehension and repeat use; proceed only if buyer value fits the data boundary.",
    },
    12: {
        6: "An iPhone habit app that turns permitted activity summaries into understandable missions and noncash points.",
        9: "3,000 / 5,000 / 7,000 eligible steps earn 10 / 20 / 30 points. Highest tier only; upgrades add the difference. Rest has no penalty.",
        11: "Email login → separate consent → read-only HealthKit → missions → ledger → privacy controls.",
        13: "Native prototype; two-iPhone tests, admin, sponsor isolation and cohort reports remain open. [E1–E3]",
    },
    13: {
        1: "User value and a testable business model",
        6: "Understand progress with less manual effort. Measure task success and matured D7 / D28 return rates.",
        8: "Adults use it free. Organisations may fund voluntary programmes; sponsors may buy generic placements.",
        10: "PROPOSED VALUE CAPTURE",
        11: "Unvalidated monthly prices: HK$15,000 per organisation; HK$10,000 per sponsor unit. [E4]",
        13: "Assumed fixed costs: HK$90,000 / month, plus user and delivery costs. At 8 clients + 1 sponsor: HK$5,000 surplus.*",
        15: "Will noncash value retain users, and will fees cover delivery capacity? The illustrative model is in Appendix 18.",
    },
    14: {
        5: "WHAT THE CURRENT EVIDENCE SUPPORTS",
        6: "Continue with noncash habit support and buyer discovery.\nGate expansion on device reliability, retention and willingness to pay.",
        12: "Users return without cash rewards.",
        13: "100 matured activations across two recruitment channels.",
        14: "D7 ≥35%\nD28 ≥20%",
        16: "Buyers fund content and support.",
        17: "5 budget interviews, then 2 authorised paid pilots after readiness gates.",
        18: "2 paid pilots;\npositive contribution",
    },
    15: {
        3: "HealthLoop · Hong Kong habit service · Team A / B / C / D",
    },
    16: {
        5: "CLAIM / SLIDES",
        6: "SOURCE ID / RECORD",
        7: "DATE AND EVIDENCE LIMIT",
        9: "Market and alternatives\nSlides 4, 7–8",
        10: "S1 · Hong Kong DH report\nS2 · Apple; S3 · Sweatcoin",
        11: "2023 survey / official product pages.\nChecked 5 Oct 2026; no demand proof.",
        13: "Data and sponsor boundary\nSlides 12–14, 19",
        14: "S4 · Apple review guidelines\nE1 · Brief and requirements",
        15: "Checked 5 Oct 2026.\nNo platform or legal approval claimed.",
        17: "Prototype and acceptance\nSlides 3, 8, 12, 19",
        18: "E2 · Status and test evidence\nE3 · Simulator walkthrough",
        19: "Sep/Oct 2026 engineering records.\nReal-device acceptance remains open.",
        21: "Business and research plan\nSlides 5–6, 9–11, 13–14, 17–18",
        22: "E4 · Existing business proposal\nProposed research in this deck",
        23: "28 Sep 2026 assumptions.\nNo customer, revenue or retention proof.",
    },
    17: {
        1: "Appendix: proposed research instrument",
        6: "PLANNED ANALYSIS PROCESS",
        7: "1. Tell me about the last time you tried to be more active.",
        9: "2. Which tool did you use, and what happened next?",
        11: "3. Where did the routine become difficult?",
        13: "4. What would make you stop using an app like this?",
        15: "5. How do you decide whether to share activity data?",
        17: "Recruit across work, language and device contexts, including sceptics and people who decline consent.",
        19: "Code trigger, workaround, friction, privacy and value. D synthesises; A checks interpretation. Link every claim to a participant code.",
        21: "Actively test whether existing tools suffice, points add no value or buyers reject the price. Preserve contrary cases; no findings are claimed yet.",
    },
}


NOTES = {
    1: """Opening, about 30 seconds. HealthLoop is an iOS-first health habit service for Hong Kong adults. We have enough technical work to discuss a concrete product, but we still need evidence that people will use it repeatedly and that a buyer will pay for operating support. The proposal is to test those questions in a bounded pilot after native and privacy readiness gates pass. The app uses Simplified Chinese; this research presentation follows the English template. The date is the evidence review date, not a promised launch or presentation date. Team letters are the project’s existing role labels, not invented member names. There is no in-app AI doctor or chatbot; Codex is a development tool. The main presentation uses the template’s 20-minute allocation, including this opening, followed by five minutes of questions.""",
    2: """Lead with the decision: continue a narrow noncash habit proposition and test repeat use before expanding incentives or building a broad commercial platform. The segment and pain statement are working hypotheses, not interview findings. The initial Hong Kong focus and iPhone platform come from the approved implementation brief. There is no market-size estimate because the available evidence does not establish the eligible, reachable or paying audience. Health reads, cloud synchronisation, marketing and reminders have separate choices. A person can decline optional choices and still access the basic app. No wallet, NFT purchase, ad interaction or insurance policy is required. Core points cannot be transferred, redeemed for cash or promised conversion to future tokens. The proposal must stand on the usefulness of the experience rather than speculative reward value.""",
    3: """This is a responsibility map, not a biography slide. The repository does not supply member names, qualifications or verified access to recruitment partners, so no credentials or institutional relationships are invented. B owns mobile and native delivery; C owns server rules, data permissions and accounting. Only B and C own the 58 technical requirements and review one another. Their human acceptance is still pending; generated code and passing tests do not confer approval. A coordinates product, commercial research and legal review, without replacing a qualified adviser. D supports research, design, recruitment and black-box QA. The concrete team asset is an existing implementation and a recorded simulator flow. Those assets make a prototype study feasible, but neither is evidence of an unmet customer need.""",
    4: """Explain the job without assuming that another health app is the answer. We want to understand how adults fit activity into ordinary workdays, what they already do, and what would count as a better outcome. Desk-based adults are a proposed first recruitment segment rather than a measured market. Initial research can compare ages 25–44 and 45–64 as exploratory strata; the product’s general audience remains adults aged 18 or older. We must check the suitability of Simplified Chinese for the Hong Kong participants we can actually reach. Traditional Chinese support has not been implemented. Apple Health is a credible existing option. We should include participants who say it is already sufficient, and people who prefer no app. The proposed service does not offer rehabilitation, diagnosis or a prescription for how much exercise someone should do.""",
    5: """Separate the completed method from the proposed method. Completed work for this deck consists of an official-source desk review and an inspection of the project brief, existing commercial proposal, implementation status and recorded technical evidence. There are no user interviews, survey responses or buyer conversations in the material reviewed. Proposed research proceeds by dependencies: approve the minimal research protocol; interview twelve adults; include eight of them in task-based usability sessions; then, after device readiness and first-party measurement are in place, observe a cohort with at least one hundred matured activations. Separately, interview five people with budget responsibility. Use recent specific behaviour before describing the concept. This sequence informs the segment, onboarding language and commercial scope rather than claiming validation from stated enthusiasm alone.""",
    6: """The count of zero refers to external primary-research participants documented in the reviewed repository. Synthetic accounts, automated test cases and an operator’s simulator walkthrough are not participants. Recruitment through two independent community or workplace channels is proposed; no access or partner permission is assumed. D can support recruitment and notes, with A coordinating scope and consent. A twelve-person qualitative sample helps discover friction but cannot estimate population prevalence. Include disconfirming participants, different levels of app familiarity and people who refuse optional cloud or marketing consent. Keep employer participation voluntary and avoid supervisor access to individual records. Obtain separate permission for recording, define storage, withdrawal and deletion before starting, and keep any real health records out of the coding session. A device study must never write invented samples into HealthKit.""",
    7: """S1 is the Hong Kong Department of Health’s Health Behaviour Survey 2023, published in 2025. The displayed statistics refer to adults aged eighteen or above. The overall survey enumerated 4,839 people aged fifteen or above from 2,145 households between July and November 2023; that overall sample is not the adult analytic denominator. Sitting or reclining excludes sleep. The data describe a population context and cannot be multiplied into HealthLoop revenue or taken as evidence that the app works. S2 confirms that Apple Health already performs basic tracking and aggregation. S3 describes Sweatcoin’s movement-based units and marketplace. These alternatives make our central test harder: does clear, localised mission feedback add enough value to be worth using? We make no comparison of clinical outcomes, retention or local merchant availability.""",
    8: """These are desk and technical findings, not completed fieldwork. Public context and established alternatives justify investigating the problem while leaving demand unproven. The 5 October application record reports a revised walk-to-points experience, server mission progress and recovery of accepted-summary claims after restart. Its logged checks passed: 355 unit/guard tests in 20 files, lint/types, 34 client-file boundaries, all 58 requirement assignments, 68 SQL groups, 25 cross-layer tests, 34 Edge/worker tests and one real local HTTP scenario with six groups. iOS export and native simulator compilation also passed. Those checks were run by the separate application session, not this proposal session. The current UI walkthrough was blocked by a locked Mac, so September’s successful recording remains evidence of the earlier interface only. No two-iPhone acceptance, retention cohort or human sign-off is established.""",
    9: """The original template asks for an evidence-based persona. At the current stage, labelling an invented person as research-derived would overstate the evidence. We therefore use a provisional segment profile and explicitly identify the assumptions to test. The goal is to learn whether people actually experience a gap between viewing records and maintaining a routine. We do not know their income, exact age, health status, shopping behaviour or emotional response. Do not add those details for realism. The first design implication is modest: a person should understand the next action, the available data and the consequence of their consent choices. Research should either replace this profile with a traceable synthesis or reject it. A non-user who says the existing tracker is enough is useful evidence, not a failed recruit.""",
    10: """Walk through this as a future scenario to test, not a witnessed customer journey. The trigger is a workday break. The person checks an existing activity record, encounters a HealthLoop mission and decides whether to allow the relevant reads and minimum cloud summary. Only a server-confirmed result becomes available points. Empty or unavailable data must remain honest; it must not be interpreted as zero activity or proven refusal of HealthKit read access. During a task session, ask the participant to explain each screen in their own words. Observe where they hesitate and whether the first action is obvious. Do not coach the person through a confusing permission decision. The outcome we want to learn is whether understandable feedback provides enough value for a later voluntary return. Feelings and friction remain hypotheses until observed.""",
    11: """The three rows separate possible pains, the strength of evidence and possible interventions. The first row has public context but no direct user confirmation. The second has real product rules and an existing interface, but no participant comprehension test. The third is a business hypothesis from the earlier proposal, with no recorded buyer interview. Prioritise research that could overturn the product decision: whether users understand the mission and whether they return when points have no cash value. Then test a buyer’s actual need for programme delivery. Do not assume that an organisation wants an employee health dashboard. If buyers require personal health profiles or health-based advertising, that demand does not fit the current product boundary. Human acceptance and commercial approval remain separate from B/C’s implementation work.""",
    12: """Explain the core flow in product language. A person signs in by email code, confirms adulthood and saves separate consent choices. On a capable device, read-only HealthKit provides available activity data; raw sleep and heart-rate series remain on-device. With cloud consent, only the minimum activity summary reaches the server. Daily tiers award the highest entitlement: 3,000 eligible steps earns 10 points; 5,000 earns 20; 7,000 earns 30. An upgrade from 10 to 20 grants only another 10. A weekly mission awards 20 once for at least three distinct days meeting the pinned goal, 3,000 by default. Rules use Asia/Hong_Kong and a next-day noon late-sync cutoff. These are product rules, not medical advice. Current reward codes are demonstrations without merchant redemption rights. No wallet, cash conversion, future-token rights or reward for ad interactions belongs to the core.""",
    13: """The existing business proposal prioritises an organisation service package: generic habit content, onboarding, participant support and administrative delivery. The organisation would buy that service, not employee health records. A separate hypothesis is clearly labelled, non-personalised sponsorship sold by placement. Sponsor events must exclude health values, health-derived groups, task outcomes, wallet identifiers and identifiers joinable to health records. These are proposed offerings: enterprise capability F03 requires separate scope approval; P23/P24 sponsorship and P30 measurement are not implemented. No contracts, customers or revenue are documented. The proposed prices and cost allocations are unvalidated inputs copied from E4. The small example surplus is not profit guidance: it excludes tax and other costs and assumes delivery capacity. Appendix 18 shows the model and adverse scenarios. Our immediate commercial task is buyer discovery, followed only by separately authorised pilots after readiness gates.""",
    14: """The evidence has narrowed the next decision, not proved product-market fit. Start with usability: at least six of eight task participants should finish the mission/ledger/privacy walkthrough unaided, with no serious misunderstanding of cash value or consent; otherwise revise and retest. This criterion is proposed in this deck. For retention, E4 defines activation as login, adult confirmation, saved choices and a home or mission view within seven days of registration; optional health, marketing and notification consent are not conditions. Set activation as D0. D7 is a voluntary core-page view on D6–8; D28 is D26–30. Include only fully matured observation windows, report sample size, uncertainty and channel differences, and exclude synthetic accounts. The 35%/20% thresholds are proposed decisions, not benchmarks or results. Buyer validation requires actual budget and procurement evidence; two authorised paid pilots with positive direct contribution are a subsequent proposed gate, not a current commitment.""",
    15: """Allow five minutes for questions. Keep the proposal’s decision clear: obtain evidence of device reliability, understandable interaction, voluntary return and a buyer who values service delivery within the privacy boundary. Refer to the appendix for model arithmetic, evidence sources and release dependencies. There is no signed commercial partner, measured retention cohort, clinical outcome or verified physical-device health flow to report. If asked about AI, explain that development uses Codex; an AI doctor, assistant subscription or chatbot is not an assumed product feature. If asked about Web3, explain that the independent synthetic-only Lab is unbuilt and does not provide a payout route for real health activity or core points. The next technical check is the updated simulator walkthrough once the Mac is unlocked, followed by the two-iPhone matrix when signing and devices are available. Admin and privacy-operation work remain independent dependencies.""",
    16: """Source register. S1–S4 are official external sources checked for this deck. E1–E4 are the repository records inspected locally. Full source locations and their limitations follow below. Slide references support their associated claims; they do not turn a proposed design into measured performance. September’s simulator recording describes the earlier interface. A concurrent 5 October application session added a newer engineering checkpoint, which this deck incorporates as recorded evidence; this proposal session did not rerun those tests. Other sessions’ application and instruction edits were preserved. Technical acceptance, legal approval, independent audit and platform approval remain open. No participant IDs or quotes are fabricated. The research instrument and additional study criteria are proposed design choices.""",
    17: """Use the interview questions before pitching the solution. Ask for a recent concrete event rather than agreement with a benefit statement. Follow up with what happened, which tool was used and what the person did instead. Obtain separate permission before a usability task or recording. The prototype task can ask participants to find the mission, explain what a point is worth, identify why data is missing, change optional consent and locate export/deletion controls without submitting destructive actions. Do not expose real participant health data to the coding environment. For buyers, separately ask about the last programme funded, budget ownership, procurement, delivery burden and the reaction to a defined HK$15,000 monthly service package. D can group notes; A checks interpretation; B/C review technical implications. Retain participant-code links and contrary cases. No interviews, coding exercise or synthesis has yet been performed.""",
    18: """All financial figures below are illustrative planning assumptions from E4, not observed revenue, quotes, approved expenditure or a forecast. Let N be paying organisations, each with 100 monthly active users (MAU), S be sponsor units, and add 500 free MAU. Organisation price is HK$15,000 per month with HK$2,500 additional delivery cost; sponsor price is HK$10,000 with HK$2,000 content/review cost. Fixed monthly allocations of HK$90,000 consist of B/C technical resources 50,000, A/D operations/content/QA 20,000, privacy/security/legal support 7,000, base cloud/tools 8,000 and contingency 5,000. These are internal allocations, not quoted full-time staffing costs. Each MAU adds HK$8 and revenue carries a 2% collection/bad-debt allowance. Result = 11,400N + 7,800S − 94,000. Avoid double-counting fixed people and additional delivery. With one sponsor, a HK$12,000 client price lowers client contribution to HK$8,460 and requires eleven clients for cost coverage. At HK$16 per MAU, the eight-client/one-sponsor case loses HK$5,400. A zero-revenue stress case with 500 MAU uses HK$94,000 monthly; three billing cycles plus an assumed HK$60,000 one-off reserve imply HK$342,000, not an approved funding ask or a development schedule. Obtain quotes and prove capacity before relying on any figure.""",
    19: """Progress by dependencies and evidence, without a promised launch date. Gate one requires actual read-only HealthKit and source behaviour on at least two compatible iPhones, signing/installability, account and consent edge cases, accessibility and the real minimum-summary-to-ledger path. Gate two requires the remaining safe operations: administrator task publishing and live MFA, retention and deletion replay after restore, sponsorship isolation before offering placements and protected cohort measurement before claiming retention. Gate three requires an approved research/pilot scope, a buyer with budget, a clear service contract, actual delivery and payment, plus reviewed unit costs. B and C remain the only technical requirement owners and cross-reviewers. A coordinates business/product/legal work; D supports research/design/QA. External provisioning, spending, publication, deployment and transaction broadcasts still need explicit operator approval. The Lab is not implemented: it must have separate identities/database, synthetic tasks and Base Sepolia chain 84532 only. It is excluded from core revenue and there is no health-to-token or mainnet path in this proposal.""",
}


SOURCE_IDS = {
    1: ["E1", "E4"], 2: ["S1", "E1", "E2", "E3"],
    3: ["E1", "E2", "E3"], 4: ["S1", "S2", "E1", "E4"],
    5: ["S1", "S2", "S3", "S4", "E4"], 6: ["E4"],
    7: ["S1", "S2", "S3"], 8: ["S1", "S2", "S3", "E2", "E3"],
    9: ["E1", "E4"], 10: ["E1", "E4"], 11: ["S1", "S2", "E1", "E3", "E4"],
    12: ["E1", "E2", "E3"], 13: ["S4", "E1", "E4"], 14: ["E1", "E4"],
    15: ["E1", "E2", "E4"], 16: list(SOURCES), 17: ["E4"],
    18: ["E4"], 19: ["E1", "E2", "E4", "S4"],
}


def set_text(shape, text, size=None):
    """Preserve the template's paragraph and run styling while changing content."""
    tf = shape.text_frame
    paragraph = tf.paragraphs[0]
    pp = deepcopy(paragraph._p.pPr) if paragraph._p.pPr is not None else None
    rp = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if p._p.pPr is not None:
            p._p.remove(p._p.pPr)
        if pp is not None:
            p._p.insert(0, deepcopy(pp))
        run = p.add_run()
        if rp is not None:
            run._r.insert(0, deepcopy(rp))
        run.text = line
        if size is not None:
            run.font.size = Pt(size)
    shape.name = re.sub(r"\s+", " ", text)[:90]


def add_text(slide, text, x, y, w, h, size=15, bold=False, color=INK):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.1
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = RGBColor.from_string(color)
    shape.name = re.sub(r"\s+", " ", text)[:90]
    return shape


def add_rule(slide, y):
    # Same thin rectangular rule as the template, with no table conversion.
    shape = slide.shapes.add_shape(1, Inches(.71), Inches(y), Inches(11.92), Inches(.01))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string("AAC0C4")
    shape.line.fill.background()


def clone_slide(prs, source):
    slide = prs.slides.add_slide(source.slide_layout)
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)
    if source._element.cSld.bg is not None:
        slide._element.cSld.insert(0, deepcopy(source._element.cSld.bg))
    for shape in source.shapes:
        slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
    return slide


def link_paragraph(shape, index, url):
    for run in shape.text_frame.paragraphs[index].runs:
        run.hyperlink.address = url


def build(template, output):
    prs = Presentation(template)
    if len(prs.slides) != 17 or "COMM7115" not in prs.slides[0].shapes[4].text:
        raise ValueError("Expected the supplied 17-slide COMM7115 research template")
    clone_slide(prs, prs.slides[15])  # 18: illustrative unit economics
    clone_slide(prs, prs.slides[6])   # 19: readiness dependencies

    for number, values in CONTENT.items():
        slide = prs.slides[number - 1]
        for index, text in values.items():
            set_text(slide.shapes[index], text)

    # Fit new content into the existing editorial grid without shrinking body copy.
    for number, indices, height in [
        (2, [6, 8], 1.10), (2, [11], 1.00), (4, [6, 8, 11, 13], 1.30),
        (6, [6, 8], 1.30), (6, [11, 13], 1.40),
        (7, [9, 10, 11, 13, 14, 15, 17, 18, 19], 1.08),
        (9, [9, 11], 1.00), (9, [13, 15], 1.02),
        (10, [8, 12, 16, 20], 1.20),
        (11, [9, 10, 11, 13, 14, 15, 17, 18, 19], 1.00),
        (12, [9, 11], 1.27), (13, [6, 8], 1.10), (13, [11, 13], 1.20),
        (17, [17], .88), (17, [19, 21], 1.14),
    ]:
        for index in indices:
            prs.slides[number - 1].shapes[index].height = Inches(height)

    # Team has four people, so replace the unused fifth-member row with whitespace.
    s = prs.slides[2]
    original_shapes = list(s.shapes)
    for indices, y in [([9, 10, 11], 2.50), ([13, 14, 15], 3.25),
                       ([17, 18, 19], 4.00), ([21, 22, 23], 4.75)]:
        for index in indices:
            original_shapes[index].top = Inches(y)
            original_shapes[index].height = Inches(.61)
    for index, y in [(12, 3.14), (16, 3.89), (20, 4.64)]:
        original_shapes[index].top = Inches(y)
    for index in [24, 25, 26, 27]:
        el = original_shapes[index]._element
        el.getparent().remove(el)
    add_text(s, "Roles are assignments, not CV claims. B/C human acceptance remains pending. [E1–E3]",
             .71, 6.55, 11.9, .25, 10.5, color=MUTED)

    # Long findings remain one visually coherent statement with an evidence line.
    for i in [6, 11, 16]:
        prs.slides[7].shapes[i].width = Inches(10.7)
    add_text(prs.slides[7], "Desk and technical findings · Primary user research remains pending.",
             .71, 6.38, 11.9, .28, 12, color=TEAL)

    add_text(prs.slides[11], "Points are nontransferable and noncash · No future-token rights, wallet requirement or ad rewards",
             .71, 3.49, 11.9, .25, 12, color=TEAL)
    add_text(prs.slides[12], "*Illustrative scenario. Enterprise and sponsor features need approval and implementation; no revenue is documented.",
             .71, 6.74, 11.9, .22, 10, color=MUTED)
    add_text(prs.slides[13], "Proposed decision rules, not results. Native and privacy readiness precede pilots. [E1, E4]",
             .71, 6.58, 11.9, .28, 11.5, color=MUTED)
    add_text(prs.slides[14], "Which evidence should decide whether HealthLoop proceeds to a pilot?",
             .71, 4.80, 11.5, .75, 21, color="B7D8DA")
    add_text(prs.slides[15], "Source IDs are used throughout. Full links and evidence scope are included in the speaker notes.",
             .71, 6.55, 11.9, .27, 11, color=MUTED)
    add_text(prs.slides[16], "Proposed instrument and synthesis process; no participant quotes or coded findings exist yet.",
             .71, 6.42, 11.9, .28, 11.5, color=MUTED)
    prs.slides[16].shapes[18].top = Inches(3.48)

    # Links are attached to the source labels while retaining template styling.
    for idx, source_id in [(10, "S1"), (14, "S2"), (18, "S3")]:
        link_paragraph(prs.slides[6].shapes[idx], 0, SOURCES[source_id][1])
    link_paragraph(prs.slides[15].shapes[10], 0, SOURCES["S1"][1])
    link_paragraph(prs.slides[15].shapes[14], 0, SOURCES["S4"][1])

    # Appendix 18 uses the template's table-like grid for a compact financial model.
    s = prs.slides[17]
    for shape in list(s.shapes)[5:]:
        shape._element.getparent().remove(shape._element)
    set_text(s.shapes[0], "APPENDIX")
    set_text(s.shapes[1], "Appendix: illustrative delivery economics")
    set_text(s.shapes[4], "18   /   BACKUP")
    add_text(s, "Unvalidated HK$ per month · Planning assumptions, not a forecast [E4]",
             .71, 1.91, 11.9, .36, 14, color=TEAL)
    columns = [(.71, 3.1), (4.20, 1.5), (6.05, 1.85), (8.27, 1.85), (10.5, 1.95)]
    for (x, w), label in zip(columns, ["CLIENTS / SPONSORS", "MAU", "REVENUE", "COST", "RESULT"]):
        add_text(s, label, x, 2.45, w, .26, 10.5, True, TEAL)
    add_rule(s, 2.80)
    scenarios = [(3, 0), (6, 1), (8, 1)]
    for j, (n, sp) in enumerate(scenarios):
        mau = 500 + 100 * n
        revenue = 15000 * n + 10000 * sp
        cost = 90000 + 8 * mau + 2500 * n + 2000 * sp + revenue * .02
        result = revenue - cost
        expected = 11400 * n + 7800 * sp - 94000
        assert result == expected
        result_text = f"+{result:,.0f}" if result > 0 else f"−{abs(result):,.0f}"
        values = [f"{n} organisations / {sp} sponsor", f"{mau:,}", f"{revenue:,}", f"{cost:,.0f}", result_text]
        for (x, w), value in zip(columns, values):
            add_text(s, value, x, 2.98 + .67 * j, w, .40, 16.5, bold=j == 2)
        add_rule(s, 3.50 + .67 * j)
    add_text(s, "MAU = 500 + 100N     Revenue = 15,000N + 10,000S", .71, 5.04, 11.9, .30, 14.5)
    add_text(s, "Cost = 90,000 + 8 × MAU + 2,500N + 2,000S + 2% × Revenue", .71, 5.42, 11.9, .30, 14.5)
    add_text(s, "Cost coverage: 9 organisations without sponsorship; 8 with one sponsor.",
             .71, 5.91, 11.9, .33, 15.5, True)
    add_text(s, "Sensitivity: HK$12,000 client price needs 11 clients + 1 sponsor; HK$16 / MAU turns the 8 + 1 case into a HK$5,400 loss.",
             .71, 6.34, 11.9, .45, 12.5)
    add_text(s, "Excludes tax, real-benefit procurement and fully costed staffing. Delivery capacity and willingness to pay remain unproven.",
             .71, 6.83, 11.9, .18, 9.5, color=MUTED)

    # Appendix 19 documents acceptance dependencies without a delivery calendar.
    s = prs.slides[18]
    updates = {
        0: "APPENDIX", 1: "Appendix: readiness gates and ownership", 4: "19   /   BACKUP",
        5: "DEPENDENCY GATE", 6: "EVIDENCE REQUIRED", 7: "RESPONSIBILITY",
        9: "1 · Native core", 10: "Two iPhones; real read → sync → claim; consent, accessibility and signing.",
        11: "B owns mobile; C reviews. Physical-device acceptance is still open.",
        13: "2 · Safe operation", 14: "Admin and MFA; deletion/restore; sponsor separation; cohort reporting.",
        15: "C owns backend; B reviews. Relevant controls must pass before use.",
        17: "3 · Commercial pilot", 18: "Approved scope; buyer budget; authorised contract, delivery and payment.",
        19: "A coordinates; D supports. B/C retain technical ownership.",
    }
    for idx, value in updates.items():
        set_text(s.shapes[idx], value)
    for idx in [9, 10, 11, 13, 14, 15, 17, 18, 19]:
        s.shapes[idx].height = Inches(1.10)
    add_text(s, "Lab remains unbuilt: separate identities/database, synthetic tasks, Base Sepolia 84532 only. No core-health-to-token path. [E1–E2]",
             .71, 6.37, 11.9, .47, 12.5, color=TEAL)

    # Notes carry presenter context, source locations, qualifiers and decision criteria.
    for number, slide in enumerate(prs.slides, 1):
        for element in slide._element.iter():
            if element.get("typeface") == "Bitstream Charter":
                element.set("typeface", FONT)
        refs = []
        for source_id in SOURCE_IDS[number]:
            title, location, limit = SOURCES[source_id]
            refs.append(f"{source_id} — {title}\n{location}\n{limit}")
        text = NOTES[number].strip() + "\n\nSources and evidence scope\n\n" + "\n\n".join(refs)
        slide.notes_slide.notes_text_frame.text = text

    props = prs.core_properties
    props.title = "HealthLoop Business Proposal Research Report"
    props.subject = "COMM7115 research proposal, evidence and validation plan"
    props.author = "HealthLoop team"
    props.last_modified_by = "HealthLoop team"
    props.keywords = "HealthLoop, business proposal, Hong Kong, noncash points"
    props.comments = "Completed from the user-supplied COMM7115 template. Evidence reviewed 5 October 2026."
    props.created = props.modified = datetime(2026, 10, 5, 0, 0, 0)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)

    # An editable plain-text checkpoint preserves the delivered content and evidence.
    md = ["# HealthLoop Business Proposal Research Report", "",
          "English presentation completed from the supplied 17-slide template, with two additional backup slides.",
          "Evidence reviewed 5 October 2026. Application records are dated 24–28 September and 5 October; this proposal session did not rerun those tests.", ""]
    for number, slide in enumerate(prs.slides, 1):
        title = slide.shapes[1].text.replace("\n", " ")
        md.extend([f"## Slide {number} {title}", "", "### Slide text", ""])
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text.strip():
                md.append(shape.text + "\n")
        md.extend(["### Speaker notes", "", slide.notes_slide.notes_text_frame.text, ""])
    output.with_suffix(".md").write_text("\n".join(md), encoding="utf-8")
    print(f"Created {len(prs.slides)} editable slides with notes: {output}")


if __name__ == "__main__":
    parser = ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=ROOT / "docs/proposals/templates/Business_Proposal_Research_Report_Template.pptx")
    parser.add_argument("--output", type=Path, default=OUT)
    args = parser.parse_args()
    build(args.template, args.output)
